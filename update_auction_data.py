import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import List, Optional

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
from playwright.sync_api import sync_playwright


AUCTION_URL = "https://www.rbauction.com/heavy-equipment-auctions"
WORKBOOK_PATH = Path("Copy of RBA Auction Data V2.xlsx")
WORKSHEET_NAME = "Quarter Summaries"
START_DATA_ROW = 6

COL_YEAR = 2
COL_MONTH = 3
COL_TAG = 4
COL_DATE = 5
COL_FORMULA_F = 6
COL_AUCTION = 8
COL_DAYS = 9
COL_LOTS = 10
COL_QTD_DAYS = 11
COL_QTD_LOTS = 12

DEBUG_SCREENSHOT = "debug_rba_page.png"
DEBUG_HTML = "debug_rba_page.html"


@dataclass
class AuctionRecord:
    auction_date_text: str
    auction_name: str
    location_text: str
    lots: int
    days: int
    source_url: Optional[str] = None

    @property
    def parsed_date(self) -> datetime:
        text = self.auction_date_text.strip()
        year = datetime.today().year

        matches = re.findall(r"([A-Z][a-z]{2})\s(\d{1,2})", text)
        if not matches:
            raise ValueError(f"Could not parse date from: {self.auction_date_text}")

        month, day = matches[0]
        return datetime.strptime(f"{month} {day} {year}", "%b %d %Y")

    @property
    def month_num(self) -> int:
        return self.parsed_date.month

    @property
    def tag_value(self) -> str:
        return f"{self.month_num}-{str(self.parsed_date.year)[-2:]}"


def parse_days(date_text: str) -> int:
    matches = re.findall(r"([A-Z][a-z]{2})\s(\d{1,2})", date_text)

    if len(matches) >= 2:
        year = datetime.today().year
        start_month, start_day = matches[0]
        end_month, end_day = matches[-1]

        start_dt = datetime.strptime(f"{start_month} {start_day} {year}", "%b %d %Y")
        end_dt = datetime.strptime(f"{end_month} {end_day} {year}", "%b %d %Y")
        return (end_dt - start_dt).days + 1

    nums = [int(x) for x in re.findall(r"\d{1,2}", date_text)]
    if len(nums) >= 2:
        return nums[-1] - nums[0] + 1

    return 1


def parse_lots(text: str) -> Optional[int]:
    match = re.search(r"(\d{1,3}(?:,\d{3})*)\s+Items?", text)
    if not match:
        return None
    return int(match.group(1).replace(",", ""))


def save_debug(page) -> None:
    try:
        page.screenshot(path=DEBUG_SCREENSHOT, full_page=True)
    except Exception:
        pass

    try:
        html = page.content()
        Path(DEBUG_HTML).write_text(html, encoding="utf-8")
    except Exception:
        pass


def dismiss_cookie_or_overlay(page) -> None:
    candidate_selectors = [
        'button:has-text("Accept")',
        'button:has-text("I Accept")',
        'button:has-text("Accept All")',
        'button:has-text("Allow All")',
        '[id*="accept"]',
        '[class*="accept"]',
    ]

    for selector in candidate_selectors:
        try:
            loc = page.locator(selector).first
            if loc.is_visible(timeout=1000):
                loc.click(timeout=1000)
                page.wait_for_timeout(1000)
                return
        except Exception:
            continue


def open_past_tab(page) -> None:
    selectors = [
        '[data-testid="sold-upcoming-toggle-past"]',
        'button[aria-label="Past"]',
        'button[value="past"]',
        'button:has-text("Past")',
        'text=Past',
    ]

    for selector in selectors:
        try:
            loc = page.locator(selector).first
            loc.wait_for(state="visible", timeout=5000)
            loc.click(timeout=5000)
            page.wait_for_timeout(2500)
            return
        except Exception:
            continue

    title = ""
    snippet = ""
    try:
        title = page.title()
    except Exception:
        pass

    try:
        body_text = page.locator("body").inner_text(timeout=2000)
        snippet = body_text[:1200]
    except Exception:
        pass

    save_debug(page)
    raise RuntimeError(
        "Could not find/click the Past tab. "
        f"Page title seen by GitHub Actions: {title!r}. "
        f"Body starts with: {snippet!r}. "
        f"Saved debug files: {DEBUG_SCREENSHOT}, {DEBUG_HTML}"
    )


def scrape_auctions() -> List[AuctionRecord]:
    records: List[AuctionRecord] = []

    with sync_playwright() as p:
        browser = p.chromium.launch(
            headless=True,
            args=[
                "--disable-blink-features=AutomationControlled",
                "--no-sandbox",
                "--disable-dev-shm-usage",
            ],
        )

        context = browser.new_context(
            user_agent=(
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/122.0.0.0 Safari/537.36"
            ),
            locale="en-US",
            viewport={"width": 1440, "height": 2200},
        )

        page = context.new_page()
        page.goto(AUCTION_URL, wait_until="domcontentloaded", timeout=60000)
        page.wait_for_timeout(6000)

        dismiss_cookie_or_overlay(page)
        open_past_tab(page)

        card_selectors = [
            '[data-testid^="auction-card-"]',
            'a[href*="/heavy-equipment-auctions/"]',
        ]

        cards = None
        card_count = 0

        for selector in card_selectors:
            try:
                cards = page.locator(selector)
                card_count = cards.count()
                if card_count > 0:
                    break
            except Exception:
                continue

        if not cards or card_count == 0:
            save_debug(page)
            raise RuntimeError(
                f"No auction cards found after opening Past tab. "
                f"Saved debug files: {DEBUG_SCREENSHOT}, {DEBUG_HTML}"
            )

        for i in range(card_count):
            card = cards.nth(i)

            try:
                text = card.inner_text(timeout=2000)
            except Exception:
                continue

            date_match = re.search(
                r"([A-Z][a-z]{2}\s\d{1,2}(?:\s-\s(?:[A-Z][a-z]{2}\s)?\d{1,2})?)",
                text,
            )
            lots = parse_lots(text)

            if not date_match or lots is None:
                continue

            auction_name = ""
            location_text = ""
            href = None

            try:
                auction_name = card.locator("h5").first.inner_text(timeout=1000).strip()
            except Exception:
                lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
                if len(lines) >= 2:
                    auction_name = lines[1]

            try:
                href = card.get_attribute("href")
            except Exception:
                href = None

            if not auction_name:
                continue

            lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
            if len(lines) >= 5:
                location_text = lines[-1]

            records.append(
                AuctionRecord(
                    auction_date_text=date_match.group(1),
                    auction_name=auction_name,
                    location_text=location_text,
                    lots=lots,
                    days=parse_days(date_match.group(1)),
                    source_url=href,
                )
            )

        browser.close()

    deduped = {}
    for record in records:
        key = (record.parsed_date.strftime("%Y-%m-%d"), record.auction_name.lower())
        deduped[key] = record

    return list(deduped.values())


class AuctionWorkbookUpdater:
    def __init__(self):
        self.wb = load_workbook(WORKBOOK_PATH)
        self.ws = self.wb[WORKSHEET_NAME]

    def existing_keys(self):
        keys = set()

        for row in range(START_DATA_ROW, self.ws.max_row + 1):
            date_val = self.ws.cell(row=row, column=COL_DATE).value
            name_val = self.ws.cell(row=row, column=COL_AUCTION).value

            if date_val and name_val:
                if isinstance(date_val, datetime):
                    date_key = date_val.strftime("%Y-%m-%d")
                else:
                    date_key = str(date_val)

                keys.add((date_key.strip(), str(name_val).strip().lower()))

        return keys

    def next_row(self):
        row = START_DATA_ROW
        while self.ws.cell(row=row, column=COL_AUCTION).value:
            row += 1
        return row

    def copy_formula(self, row: int, col: int):
        if row <= START_DATA_ROW:
            return

        src = self.ws.cell(row=row - 1, column=col)
        if isinstance(src.value, str) and src.value.startswith("="):
            self.ws.cell(row=row, column=col).value = src.value

    def update(self, records: List[AuctionRecord]):
        existing = self.existing_keys()
        added = 0

        for record in sorted(records, key=lambda r: (r.parsed_date, r.auction_name)):
            key = (record.parsed_date.strftime("%Y-%m-%d"), record.auction_name.lower())
            if key in existing:
                continue

            row = self.next_row()

            self.ws.cell(row=row, column=COL_YEAR).value = record.parsed_date.year
            self.ws.cell(row=row, column=COL_MONTH).value = record.month_num
            self.ws.cell(row=row, column=COL_TAG).value = record.tag_value
            self.ws.cell(row=row, column=COL_DATE).value = record.parsed_date
            self.ws.cell(row=row, column=COL_AUCTION).value = record.auction_name
            self.ws.cell(row=row, column=COL_DAYS).value = record.days
            self.ws.cell(row=row, column=COL_LOTS).value = record.lots

            self.copy_formula(row, COL_FORMULA_F)
            self.copy_formula(row, COL_QTD_DAYS)
            self.copy_formula(row, COL_QTD_LOTS)

            fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
            for col in range(COL_YEAR, COL_QTD_LOTS + 1):
                self.ws.cell(row=row, column=col).fill = fill

            existing.add(key)
            added += 1

        self.wb.save(WORKBOOK_PATH)
        return added


def main():
    records = scrape_auctions()
    print(f"Scraped {len(records)} auctions")

    updater = AuctionWorkbookUpdater()
    added = updater.update(records)

    print(f"Rows added: {added}")


if __name__ == "__main__":
    main()
